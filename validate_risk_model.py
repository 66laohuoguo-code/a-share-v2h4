"""Validate and summarize a built weekly A-share risk model."""

from __future__ import annotations

import argparse
from datetime import datetime
import json
import math
from pathlib import Path
import sqlite3

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_DATABASE = (
    PROJECT_ROOT / "data" / "processed" / "csmar_risk_model_v1.sqlite"
)
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "outputs" / "risk_model_validation"


def json_ready(value):
    if isinstance(value, dict):
        return {key: json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_ready(item) for item in value]
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def table_exists(conn, table):
    return (
        conn.execute(
            """
            SELECT 1 FROM sqlite_master
            WHERE type='table' AND name=?
            """,
            (table,),
        ).fetchone()
        is not None
    )


def table_stats(conn, table, date_column=None):
    if not table_exists(conn, table):
        return {"exists": False}
    result = {
        "exists": True,
        "rows": conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0],
    }
    if date_column:
        minimum, maximum = conn.execute(
            f"SELECT MIN({date_column}), MAX({date_column}) FROM {table}"
        ).fetchone()
        result["min_date"] = minimum
        result["max_date"] = maximum
    return result


def latest_covariance(conn):
    model_date = conn.execute(
        "SELECT MAX(model_date) FROM weekly_factor_covariance"
    ).fetchone()[0]
    if not model_date:
        return model_date, pd.DataFrame()
    rows = pd.read_sql_query(
        """
        SELECT factor_1, factor_2, covariance
        FROM weekly_factor_covariance
        WHERE model_date=?
        """,
        conn,
        params=(model_date,),
    )
    matrix = rows.pivot(
        index="factor_1", columns="factor_2", values="covariance"
    )
    matrix = matrix.reindex(index=sorted(matrix.index), columns=sorted(matrix.columns))
    return model_date, matrix


def factor_summary(conn):
    frame = pd.read_sql_query(
        """
        SELECT model_date, factor_name, factor_return
        FROM weekly_factor_return
        ORDER BY model_date, factor_name
        """,
        conn,
    )
    if frame.empty:
        return pd.DataFrame()
    summary = (
        frame.groupby("factor_name")["factor_return"]
        .agg(["count", "mean", "std", "min", "max"])
        .reset_index()
    )
    summary["annualized_return"] = summary["mean"] * 52.0
    summary["annualized_volatility"] = summary["std"] * math.sqrt(52.0)
    summary["return_to_risk"] = (
        summary["annualized_return"] / summary["annualized_volatility"]
    )
    return summary.sort_values("factor_name").reset_index(drop=True)


def latest_specific_summary(conn):
    model_date = conn.execute(
        "SELECT MAX(model_date) FROM weekly_specific_risk"
    ).fetchone()[0]
    if not model_date:
        return model_date, pd.DataFrame()
    frame = pd.read_sql_query(
        """
        SELECT code, industry_group, specific_variance,
               specific_volatility, specific_observations
        FROM weekly_specific_risk
        WHERE model_date=?
        ORDER BY specific_volatility DESC
        """,
        conn,
        params=(model_date,),
    )
    return model_date, frame


def no_lookahead_audit(conn):
    if not table_exists(conn, "weekly_exposure"):
        return {}
    row = conn.execute(
        """
        SELECT
            SUM(observation_date > model_date),
            SUM(
                financial_available_date IS NOT NULL
                AND financial_available_date >= model_date
            ),
            SUM(
                industry_implement_date IS NOT NULL
                AND industry_implement_date > model_date
            )
        FROM weekly_exposure
        """
    ).fetchone()
    return {
        "future_market_observations": int(row[0] or 0),
        "future_financial_observations": int(row[1] or 0),
        "future_industry_observations": int(row[2] or 0),
        "passed": not any(row),
    }


def configured_model_end_date(conn):
    if not table_exists(conn, "risk_model_metadata"):
        return None
    row = conn.execute(
        "SELECT value FROM risk_model_metadata WHERE key='raw_end'"
    ).fetchone()
    if row:
        try:
            return str(json.loads(row[0]))
        except (TypeError, ValueError):
            return str(row[0])
    row = conn.execute(
        "SELECT value FROM risk_model_metadata WHERE key='config'"
    ).fetchone()
    if not row:
        return None
    try:
        return json.loads(row[0]).get("model_end_date")
    except (AttributeError, TypeError, ValueError):
        return None


def append_frame_sheet(workbook, title, frame):
    sheet = workbook.create_sheet(title)
    if frame is None or frame.empty:
        sheet.append(["No data"])
        return
    sheet.append(list(frame.columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in frame.itertuples(index=False, name=None):
        sheet.append(
            [
                value.item() if isinstance(value, np.generic) else value
                for value in row
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, column in enumerate(frame.columns, start=1):
        sample = [str(column)]
        sample.extend(
            str(value)
            for value in frame.iloc[:200, column_index - 1].dropna().tolist()
        )
        width = min(max(len(value) for value in sample) + 2, 28)
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def write_outputs(output_dir, summary, factor_frame, diagnostics, specific, covariance):
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_path = output_dir / f"risk_model_validation_{stamp}.json"
    xlsx_path = output_dir / f"risk_model_validation_{stamp}.xlsx"
    with json_path.open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=False, indent=2)

    workbook = Workbook()
    workbook.remove(workbook.active)
    summary_rows = []
    for key, value in summary.items():
        if isinstance(value, (dict, list)):
            display = json.dumps(value, ensure_ascii=False)
        else:
            display = value
        summary_rows.append({"item": key, "value": display})
    append_frame_sheet(workbook, "summary", pd.DataFrame(summary_rows))
    append_frame_sheet(workbook, "factor_returns", factor_frame)
    append_frame_sheet(workbook, "diagnostics", diagnostics)
    append_frame_sheet(workbook, "specific_risk", specific)
    if covariance is not None and not covariance.empty:
        covariance_frame = covariance.reset_index().rename(
            columns={"factor_1": "factor"}
        )
    else:
        covariance_frame = pd.DataFrame()
    append_frame_sheet(workbook, "factor_covariance", covariance_frame)
    workbook.save(xlsx_path)
    return json_path, xlsx_path


def run(args):
    conn = sqlite3.connect(args.database)
    try:
        required = {
            "stock_market_cap": "trade_date",
            "financial_pit": "available_date",
            "industry_history": "implement_date",
            "weekly_raw_exposure": "model_date",
            "weekly_exposure": "model_date",
            "weekly_factor_return": "model_date",
            "weekly_specific_return": "model_date",
            "weekly_factor_covariance": "model_date",
            "weekly_specific_risk": "model_date",
            "weekly_risk_diagnostics": "model_date",
        }
        tables = {
            table: table_stats(conn, table, date_column)
            for table, date_column in required.items()
        }
        covariance_date, covariance = latest_covariance(conn)
        specific_date, specific = latest_specific_summary(conn)
        factors = factor_summary(conn)
        diagnostics = pd.read_sql_query(
            """
            SELECT *
            FROM weekly_risk_diagnostics
            ORDER BY model_date
            """,
            conn,
        )
        eigenvalues = (
            np.linalg.eigvalsh(covariance.to_numpy(dtype=float))
            if not covariance.empty
            else np.array([])
        )
        latest_diagnostic = (
            diagnostics.iloc[-1].to_dict() if not diagnostics.empty else {}
        )
        configured_end = configured_model_end_date(conn)
        latest_model_date = latest_diagnostic.get("model_date")
        summary = {
            "database": str(args.database),
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "tables": tables,
            "no_lookahead": no_lookahead_audit(conn),
            "configured_model_end_date": configured_end,
            "latest_model_date": latest_model_date,
            "model_complete_through_configured_end": (
                latest_model_date >= configured_end
                if latest_model_date and configured_end
                else None
            ),
            "latest_universe_count": latest_diagnostic.get("universe_count"),
            "latest_financial_book_coverage": latest_diagnostic.get(
                "financial_book_coverage"
            ),
            "latest_financial_ttm_coverage": latest_diagnostic.get(
                "financial_ttm_coverage"
            ),
            "latest_industry_coverage": latest_diagnostic.get(
                "industry_coverage"
            ),
            "latest_covariance_date": covariance_date,
            "latest_covariance_factors": len(covariance),
            "latest_covariance_min_eigenvalue": (
                float(eigenvalues.min()) if len(eigenvalues) else None
            ),
            "latest_covariance_is_psd": (
                bool(eigenvalues.min() >= -1e-10) if len(eigenvalues) else None
            ),
            "latest_specific_risk_date": specific_date,
            "latest_specific_risk_count": len(specific),
        }
        summary = json_ready(summary)
        json_path, xlsx_path = write_outputs(
            args.output_dir,
            summary,
            factors,
            diagnostics,
            specific,
            covariance,
        )
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        print(f"Summary JSON: {json_path}")
        print(f"Validation workbook: {xlsx_path}")
        if not summary["no_lookahead"].get("passed", False):
            raise ValueError("The no-lookahead audit failed")
        if summary["latest_covariance_is_psd"] is False:
            raise ValueError("The latest factor covariance is not PSD")
        if (
            summary["model_complete_through_configured_end"] is False
            and not args.allow_partial
        ):
            raise ValueError(
                "The risk model is incomplete: latest model date "
                f"{summary['latest_model_date']} is earlier than configured "
                f"end date {summary['configured_model_end_date']}."
            )
    finally:
        conn.close()


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Validate and summarize the weekly A-share risk model."
    )
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument(
        "--allow-partial",
        action="store_true",
        help="Write validation outputs without failing an incomplete end date.",
    )
    args = parser.parse_args(argv)
    args.database = args.database.resolve()
    args.output_dir = args.output_dir.resolve()
    if not args.database.exists():
        parser.error(f"Database does not exist: {args.database}")
    return args


if __name__ == "__main__":
    run(parse_args())
